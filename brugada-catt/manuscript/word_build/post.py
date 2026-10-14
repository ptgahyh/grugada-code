import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

d = Document(sys.argv[1])
no_line_numbers_later = []

# ---- page setup: A4, 2.5 cm margins, continuous line numbers, page number footer
for s in d.sections:
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, m, Cm(2.5))
    s.header_distance = Cm(1.25); s.footer_distance = Cm(1.25); s.gutter = Cm(0)
    sp = s._sectPr
    ln = OxmlElement("w:lnNumType"); ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); ln.set(qn("w:distance"), "283")
    pg = sp.find(qn("w:pgMar"))
    pg.addnext(ln) if pg is not None else sp.append(ln)
    fp = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); r._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)
    r.font.size = Pt(10); r.font.name = "Times New Roman"
    no_line_numbers_later.append(fp)


PPR_AFTER = ["pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct",
             "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
             "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl",
             "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]


def no_line_numbers(p):
    ppr = p._p.get_or_add_pPr()
    if ppr.find(qn("w:suppressLineNumbers")) is not None:
        return
    el = OxmlElement("w:suppressLineNumbers")
    for child in ppr:
        if child.tag in [qn("w:" + t) for t in PPR_AFTER]:
            child.addprevious(el)
            return
    ppr.append(el)


TC_AFTER = ["shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]


def border(el, edges, sz=8):
    tcpr = el.get_or_add_tcPr()
    b = tcpr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        for child in tcpr:
            if child.tag in [qn("w:" + t) for t in TC_AFTER]:
                child.addprevious(b); break
        else:
            tcpr.append(b)
    for edge in edges:
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), "000000")
        b.append(e)


# ---- tables: three-line style, 9.5 pt, single spacing
for t in d.tables:
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    w = tblPr.find(qn("w:tblW"))
    if w is None:
        w = OxmlElement("w:tblW"); tblPr.append(w)
    w.set(qn("w:type"), "pct"); w.set(qn("w:w"), "5000")
    rows = t.rows
    for i, row in enumerate(rows):
        for cell in row.cells:
            edges = []
            if i == 0:
                edges += ["top", "bottom"]
            if i == len(rows) - 1:
                edges += ["bottom"]
            if edges:
                border(cell._tc, edges, 10 if i == 0 and "top" in edges else 6)
            for p in cell.paragraphs:
                pf = p.paragraph_format
                pf.line_spacing = 1.0; pf.space_after = Pt(1); pf.space_before = Pt(1)
                no_line_numbers(p)
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    if i == 0:
                        r.font.bold = True

# ---- column widths proportional to content (16 cm text width)
TOTAL = 9071
for t in d.tables:
    ncol = len(t.columns)
    lens = [0] * ncol
    words = [0] * ncol
    for row in t.rows:
        for j, cell in enumerate(row.cells[:ncol]):
            txt = cell.text.strip()
            longest_word = max((sum(1.45 if c.isupper() else 1 for c in w) for w in txt.split()), default=0)
            lens[j] = max(lens[j], min(len(txt), 30))
            words[j] = max(words[j], longest_word + 1)
    small = ncol >= 8
    cpx = 90 if small else 100          # dxa per character (Times, 8.5 / 9.5 pt)
    mins = [w * cpx + 230 for w in words]
    if sum(mins) > TOTAL:
        widths = [int(m * TOTAL / sum(mins)) for m in mins]
    else:
        extra = [max(l * cpx + 230 - m, 0) for l, m in zip(lens, mins)]
        spare = TOTAL - sum(mins)
        if sum(extra) == 0:
            widths = [int(m + spare / ncol) for m in mins]
        else:
            share = [min(e, spare * e / sum(extra)) for e in extra]
            left = spare - sum(share)
            widths = [int(m + sh + left * (m + sh) / (sum(mins) + sum(share))) for m, sh in zip(mins, share)]

    if small:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8.5)
    # keep each table on one page
    for row in t.rows[:-1]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True
    tblPr = t._tbl.tblPr
    tw = tblPr.find(qn("w:tblW")); tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(TOTAL))
    lay = tblPr.find(qn("w:tblLayout"))
    if lay is None:
        lay = OxmlElement("w:tblLayout")
        anchor = None
        for tag in ("w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"):
            anchor = tblPr.find(qn(tag))
            if anchor is not None:
                break
        anchor.addprevious(lay) if anchor is not None else tblPr.append(lay)
    lay.set(qn("w:type"), "fixed")
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for row in t.rows:
        for cell, w in zip(row.cells[:ncol], widths):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW"); tcpr.insert(0, tcw)
            tcw.set(qn("w:type"), "dxa"); tcw.set(qn("w:w"), str(w))

# ---- breathing room after each table
from docx.text.paragraph import Paragraph
for t in d.tables:
    nxt = t._tbl.getnext()
    if nxt is not None and nxt.tag == qn("w:p"):
        Paragraph(nxt, d._body).paragraph_format.space_before = Pt(10)

# ---- captions, figures, equations
cap = re.compile(r"^(Table|Figure) \d+\.")
paras = d.paragraphs
for k, p in enumerate(paras):
    txt = p.text.strip()
    if p._p.findall(".//" + qn("w:drawing")):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        no_line_numbers(p)
        continue
    if p._p.findall(".//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathPara"):
        no_line_numbers(p)
        continue
    if cap.match(txt):
        pf = p.paragraph_format
        pf.line_spacing = 1.15; pf.space_before = Pt(6); pf.space_after = Pt(8)
        pf.keep_together = True
        if txt.startswith("Table"):
            pf.keep_with_next = True
        for r in p.runs:
            r.font.size = Pt(10)

for ppr in d.element.body.iter(qn("w:pPr")):
    ps = ppr.find(qn("w:pStyle"))
    if ps is not None and ppr.index(ps) != 0:
        ppr.remove(ps); ppr.insert(0, ps)
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
for dpr in d.element.body.iter(M + "dPr"):
    order = ["begChr", "sepChr", "endChr", "grow", "shp", "ctrlPr"]
    kids = sorted(list(dpr), key=lambda e: order.index(e.tag.replace(M, "")) if e.tag.replace(M, "") in order else 99)
    for k in kids:
        dpr.remove(k); dpr.append(k)
for rpr in d.element.body.iter(M + "rPr"):
    tags = [e.tag.replace(M, "") for e in rpr]
    if "nor" in tags:
        for e in list(rpr):
            if e.tag in (M + "sty", M + "scr"):
                rpr.remove(e)
    order = ["lit", "nor", "scr", "sty", "brk", "aln"]
    kids = sorted(list(rpr), key=lambda e: order.index(e.tag.replace(M, "")) if e.tag.replace(M, "") in order else 99)
    for k in kids:
        rpr.remove(k); rpr.append(k)
st = d.settings.element
z = st.find(qn("w:zoom"))
if z is not None:
    st.remove(z)
    anchor = None
    for t in ("w:view", "w:writeProtection"):
        anchor = st.find(qn(t))
        if anchor is not None:
            break
    anchor.addnext(z) if anchor is not None else st.insert(0, z)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers"] + PPR_AFTER
STYLE_ORDER = ["name", "aliases", "basedOn", "next", "link", "autoRedefine", "hidden", "uiPriority", "semiHidden",
               "unhideWhenUsed", "qFormat", "locked", "personal", "personalCompose", "personalReply", "rsid",
               "pPr", "rPr", "tblPr", "trPr", "tcPr", "tblStylePr"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders"] + TC_AFTER


def reorder(el, order):
    loc = lambda e: e.tag.split("}")[-1]
    kids = sorted(list(el), key=lambda e: order.index(loc(e)) if loc(e) in order else 999)
    for k in kids:
        el.remove(k); el.append(k)


root = d.styles.element
for st_el in root.iter(W + "style"):
    reorder(st_el, STYLE_ORDER)
for ppr in root.iter(W + "pPr"):
    reorder(ppr, PPR_ORDER)
for tcpr in root.iter(W + "tcPr"):
    reorder(tcpr, TCPR_ORDER)
try:
    num = d.part.numbering_part.element
    for n in num.iter(W + "nsid"):
        n.set(W + "val", n.get(W + "val").zfill(8))
except Exception:
    pass
es = st.find(qn("w:embedSystemFonts"))
if es is not None:
    st.remove(es)
    anchor = st.find(qn("w:embedTrueTypeFonts")) or st.find(qn("w:zoom"))
    anchor.addnext(es) if anchor is not None else st.insert(0, es)
SET_ORDER = ["writeProtection", "view", "zoom", "removePersonalInformation", "removeDateAndTime", "doNotDisplayPageBoundaries",
             "displayBackgroundShape", "printPostScriptOverText", "printFractionalCharacterWidth", "printFormsData",
             "embedTrueTypeFonts", "embedSystemFonts", "saveSubsetFonts", "saveFormsData", "mirrorMargins",
             "alignBordersAndEdges", "bordersDoNotSurroundHeader", "bordersDoNotSurroundFooter", "gutterAtTop",
             "hideSpellingErrors", "hideGrammaticalErrors", "activeWritingStyle", "proofState", "formsDesign",
             "attachedTemplate", "linkStyles", "stylePaneFormatFilter", "stylePaneSortMethod", "documentType",
             "mailMerge", "revisionView", "trackRevisions", "doNotTrackMoves", "doNotTrackFormatting",
             "documentProtection", "autoFormatOverride", "styleLockTheme", "styleLockQFSet", "defaultTabStop",
             "autoHyphenation", "consecutiveHyphenLimit", "hyphenationZone", "doNotHyphenateCaps",
             "showEnvelope", "summaryLength", "clickAndTypeStyle", "defaultTableStyle", "evenAndOddHeaders",
             "bookFoldRevPrinting", "bookFoldPrinting", "bookFoldPrintingSheets", "drawingGridHorizontalSpacing",
             "drawingGridVerticalSpacing", "displayHorizontalDrawingGridEvery", "displayVerticalDrawingGridEvery",
             "doNotUseMarginsForDrawingGridOrigin", "drawingGridHorizontalOrigin", "drawingGridVerticalOrigin",
             "doNotShadeFormData", "noPunctuationKerning", "characterSpacingControl", "printTwoOnOne",
             "strictFirstAndLastChars", "noLineBreaksAfter", "noLineBreaksBefore", "savePreviewPicture",
             "doNotValidateAgainstSchema", "saveInvalidXml", "ignoreMixedContent", "alwaysShowPlaceholderText",
             "doNotDemarcateInvalidXml", "saveXmlDataOnly", "useXSLTWhenSaving", "saveThroughXslt",
             "showXMLTags", "alwaysMergeEmptyNamespace", "updateFields", "hdrShapeDefaults", "footnotePr",
             "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema", "themeFontLang",
             "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
             "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults",
             "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"]
reorder(st, SET_ORDER)
for fp in no_line_numbers_later:
    no_line_numbers(fp)
d.core_properties.title = "Concept-anchored threshold transfer for Brugada ECG classifiers — Methods and Results"
d.core_properties.author = ""
d.save(sys.argv[2])
